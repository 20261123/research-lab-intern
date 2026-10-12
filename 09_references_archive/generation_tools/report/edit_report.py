from pathlib import Path
import json, copy, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SRC=Path(r'C:\Users\user\Desktop\연구실인턴\2일차_주파수_대역폭_비교by동욱.docx')
DEST=ROOT/'2일차_주파수_대역폭_비교_labgemini4_수정본.docx'
baseline=hashlib.sha256(SRC.read_bytes()).hexdigest()
d=Document(SRC)
data=json.loads((ROOT/'stats.json').read_text())
stats=data['stats']
p=d.paragraphs

def replace(p,text,bold=None):
    props=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if props is not None:r._r.insert(0,props)
    if bold is not None:r.bold=bold

replace(p[0],'2일차 주파수 대역폭 비교 결과')
replace(p[1],'대상  세종대 기지국·건물 배치  |  실행 코드  labgemini4.m')
replace(p[2],'공통 설정  섹터당 43 dBm · NF 7 dB · 간섭 부하율 0.5\n평가 영역  x −150~150 m, y 0~550 m · 간격 2.5 m · 26,741개 위치')
replace(p[3],'실험 결과')
replace(p[4],'1 주파수 비교 대역폭 100 MHz 고정')
replace(p[5],'3.5 GHz → 28 GHz')
replace(p[6],'같은 대역폭에서는 3.5 GHz의 세 전송률 지표가 더 높았다.',True)
replace(p[7],'주파수가 높아지면 신호와 간섭이 모두 줄어든다. 이 코드에서는 잡음의 상대적 영향이 커져 SINR이 낮아진다.')
replace(p[8],'2 대역폭 비교 중심주파수 28 GHz 고정')
replace(p[10],'대역폭 확대 시 평균·중앙값·5백분위가 모두 증가했다.',True)
replace(p[11],'대역폭이 8배가 되면 총 잡음전력도 9.03 dB 증가한다.\n따라서 SINR은 낮아지며, 전송률 증가율은 8배보다 작다.')
replace(p[12],'3 두 대표 조건 비교')
replace(p[14],'이번 배치와 부하율에서는 28 GHz/800 MHz의 세 지표가 가장 높았다.',True)
replace(p[15],'평균 5.16배, 중앙값 3.73배, 5백분위 2.72배이다. 주파수와 대역폭을 함께 바꾼 결과이며, 실제 다운로드 속도를 뜻하지 않는다.')

def cell_text(cell,text):
    replace(cell.paragraphs[0],text)
for t in d.tables:
    cell_text(t.cell(0,len(t.columns)-1),'5백분위')
for i,row in enumerate(stats,1):
    for j,val in enumerate(row,1):cell_text(d.tables[0].cell(i,j),f'{val:.2f} Mbps')
for j,val in enumerate(data['freq_decrease']):cell_text(d.tables[1].cell(1,j),f'{val:.1f}% 감소')
for ti,key in [(2,'bw_ratio'),(3,'representative_ratio')]:
    for j,val in enumerate(data[key]):cell_text(d.tables[ti].cell(1,j),f'{val:.2f}배 증가')
# Preserve the source's typography and page system; remove inherited title rule.
title=d.styles['Title']
if title.element.pPr is not None:
    for e in list(title.element.pPr):
        if e.tag==qn('w:pBdr'):title.element.pPr.remove(e)
for name in ['Title','Heading 1','Heading 2']:
    s=d.styles[name];s.font.color.rgb=RGBColor(0,0,0)
    s.font.name='맑은 고딕'
    rpr=s.element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for a in list(fonts.attrib):
            if 'Theme' in a:del fonts.attrib[a]
        fonts.set(qn('w:eastAsia'),'맑은 고딕')
for t in d.tables:
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)

def para(text,bold=False):
    p=d.add_paragraph();r=p.add_run(text);r.bold=bold
    return p
def head(text):return d.add_paragraph(text,'Heading 1')
def issue(title,text):
    p=para('');p.paragraph_format.space_before=Pt(8)
    p.add_run(title+'  ').bold=True;p.add_run(text)

d.add_page_break()
head('기존 코드의 문제점과 수정사항')
para('labgemini2.m에서 발견한 계산·표시 문제를 labgemini4.m에 반영했다. 평가 영역, 간섭 조건, 통계 지표도 달라졌으므로 기존 수치 대신 수정 코드의 실행 결과로 표와 해석을 갱신했다.')
issue('1 히트맵 색상','기존 scatter에 RGB 색상 배열을 직접 넣으면 실제 SINR과 색상축의 대응이 맞지 않는다. 수정 코드에서는 SINR_dB(:,c)를 색상 데이터로 전달하고 두 그림에 caxis([-15 25])를 동일하게 적용했다. 건물 외곽선도 표시했다.')
issue('2 기지국 좌표 중복','(36.1, 418.1) m가 교내 서비스·외부 간섭 목록에 중복되어 동일 사이트를 이중 집계했다. 외부 목록에서 제거해 교내 9개, 외부 6개 사이트를 사용한다. intersect는 목록 간 중복을 경고하며, 새 중복을 자동 삭제하지는 않는다.')
issue('3 평가 영역','기존 y 상한 500 m 밖에 교내 사이트 (−126.2, 521.6) m가 있었다. 상한을 550 m로 늘리고 meshgrid로 세 조건에 공통인 26,741개 위치를 생성했다. 기지국은 평가 영역 밖에 있어도 전력 계산에는 포함되므로, 이 수정은 북쪽 사용자 영역을 추가한 것이다.')
issue('4 간섭 부하율','기존 모델은 모든 비서빙 섹터의 동일 주파수 전력을 합산했다. 수정 코드는 I = load_factor × I_all로 두고 기본값을 0.5로 설정한다. 1.0이면 전체 간섭 전력을 사용한다. 부하율 변경은 모델 가정 변경이며 실제 기지국 부하를 측정한 값은 아니다.')
issue('5 안테나 상대각','기존 x 방향 조건 분기와 정규화되지 않은 각도는 기지국 좌표를 경계로 인위적인 불연속을 만들 수 있었다. sectorPower에서 angleDiff = mod(theta − sectorAz + 180, 360) − 180으로 상대각을 [−180°, 180°) 범위로 정리했다. 건물 마스크 경계에서 나타나는 손실 변화는 별개이다.')
issue('6 하위 전송률 지표','기존 보고서의 하위 5% 평균을 5백분위로 바꿨다. 수정 코드 방식대로 오름차순 전송률의 ceil(26,741 × 0.05) = 1,338번째 값을 사용한다. 낮은 전송률 사용자들의 평균과는 다른 지표이다.')
head('계산 기준과 해석 범위')
para('잡음 N_dBm = −174 + 10 log₁₀(B_Hz) + NF이며, 100 MHz는 −87 dBm, 800 MHz는 −77.97 dBm이다. 신호·간섭·잡음을 mW로 통일한 뒤 SINR = S/(I+N), C = B log₂(1+SINR)로 계산한다. 용량 식에는 선형 SINR을 사용한다.')
para('이 구현은 모든 링크에 동일한 UMa NLOS 형태의 경로손실 식을 적용하며 LOS 확률, shadow fading, 실내 이동거리 손실은 포함하지 않는다. 건물 내부 위치에는 서비스·외부 링크 모두 low-loss 외벽 손실을 적용한다. 부하율은 간섭만 일정 비율로 줄이는 단순화이다. 결과는 위치별 단일 스트림 Shannon 전송률이며 셀 총 처리량이 아니다.')

d.add_page_break()
head('수정 코드의 히트맵과 전송률 분포')
p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'heatmaps.png'),width=Cm(17.35))
para('그림 1  100 MHz에서의 주파수 비교. 삼각형은 교내 서비스 사이트, ×는 외부 간섭 사이트, 검은 사각형은 건물이다. SINR 0 dB 미만 위치 비율은 3.5 GHz 37.78%, 28 GHz 48.69%이다.')
para('색상 범위는 −15~25 dB로 동일하다. 범위 밖 값은 끝 색상으로 포화되어 표시된다. 그림의 x 표시 범위는 −210~160 m이고 실제 평가 격자는 −150~150 m이므로, 왼쪽 빈 영역은 미평가 영역이다.')
p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'capacity_cdf.png'),width=Cm(12.8))
para('그림 2  3.5 GHz/100 MHz와 28 GHz/800 MHz의 Capacity CDF. 전체 26,741개 위치를 같은 가중치로 집계하며 CDF는 1까지 도달한다. 28 GHz/800 MHz의 평균·중앙값·5백분위가 모두 높지만, 모든 위치의 우세를 뜻하지는 않는다.')
para('실행 기록  MATLAB R2025a에서 labgemini4.m 원본을 실행했다. 난수나 반복 시행은 없는 결정론적 계산이다. 세 조건 모두 같은 사용자 격자와 기지국 배치를 사용했고, 표의 28 GHz/100 MHz 통계는 코드가 계산한 Capacity(:,2)에서 같은 방식으로 집계했다.')

d.save(DEST)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==baseline
print(DEST)
